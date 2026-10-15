import json
import logging
import os
from decimal import Decimal, InvalidOperation
from pathlib import Path
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


def amount(value):
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite() or number < 0 or number != number.quantize(Decimal('.01')):
            raise ValueError('金额必须为非负数且最多两位小数')
        return int(number * 100)
    except InvalidOperation as exc:
        raise ValueError('金额格式错误') from exc


def integer(value, positive=False):
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite() or number != number.to_integral_value() or number < (1 if positive else 0):
            raise ValueError('数量须为有效整数')
        return int(number)
    except InvalidOperation as exc:
        raise ValueError('数量格式错误') from exc


def read_table(path, required):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding='utf-8-sig')
    elif path.suffix.lower() == '.xlsx':
        frame = pd.read_excel(path, dtype=str, keep_default_na=False, engine='openpyxl')
    else:
        raise ValueError(f'不支持的文件格式: {path.suffix}')
    missing = set(required) - set(frame.columns)
    if missing:
        raise ValueError(f'{path.name} 缺少字段: {sorted(missing)}')
    return frame.map(lambda x: str(x).strip()).to_dict('records')


def logger_for(name, folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger(f'{name}.{folder.resolve()}')
    logger.setLevel(logging.INFO)
    if not logger.handlers:
        handler = logging.FileHandler(folder / 'run.log', encoding='utf-8')
        handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'))
        logger.addHandler(handler)
    return logger


def write_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.tmp')
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    os.replace(temp, path)


def write_workbook(path, sheets):
    """结果值工作簿；计算由受测 Python 完成，源数据修改后须重新运行。"""
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    for title, (fields, rows) in sheets.items():
        sheet = book.create_sheet(title)
        sheet.append(fields)
        for row in rows:
            values = []
            for key in fields:
                value = row.get(key, '')
                if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
                    value = "'" + value
                values.append(value)
            sheet.append(values)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='243B53')
            cell.font = Font(name='Microsoft YaHei', color='FFFFFF', bold=True, size=11)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Microsoft YaHei', size=11)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
            sheet.row_dimensions[row[0].row].height = 32
        for column in sheet.columns:
            letter = column[0].column_letter
            max_width = max((len(str(c.value or '')) + sum(ord(ch) > 127 for ch in str(c.value or '')) for c in column), default=12)
            sheet.column_dimensions[letter].width = min(58, max(14, max_width + 3))
        sheet.row_dimensions[1].height = 28
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.stem + '.tmp.xlsx')
    book.save(temp)
    os.replace(temp, path)
