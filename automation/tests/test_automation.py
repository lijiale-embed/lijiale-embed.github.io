import json
import sqlite3
import sys
import threading
from collections import Counter
from pathlib import Path
import pandas as pd
import pytest
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from make_samples import make
from shared import amount,integer,write_workbook,read_table
from ops_daily.pipeline import run as daily
from catalog_bot.server import make_server,validate
from catalog_bot.robot import run as robot


@pytest.fixture
def sample(tmp_path):
    make(tmp_path)
    return tmp_path


def test_daily_golden_metrics_and_exception_ids(sample):
    folder=sample/'ops_daily'
    result=daily(folder/'data',folder/'outputs')
    expected=json.loads((folder/'data/expected.json').read_text())
    assert {k:result[k] for k in expected}==expected
    issues=json.loads((folder/'outputs/exceptions.json').read_text(encoding='utf-8'))
    assert {r['record_id'] for r in issues}=={'O009','O010','O011','O012','O013','O014','O015','R004','R005','R006','R008'}
    assert sum(r['code']=='CONFLICT_DUPLICATE' for r in issues)==2
    expected_issues=Counter([('O015','CONFLICT_DUPLICATE')]*2+
        [(f'O{index:03}','INVALID_ORDER') for index in range(9,15)]+
        [(rid,'INVALID_REFUND') for rid in ['R004','R005','R008']]+[('R006','REFUND_EXCEEDS_ORDER')])
    assert Counter((r['record_id'],r['code']) for r in issues)==expected_issues


def test_daily_repeat_is_stable_and_workbook_matches(sample):
    folder=sample/'ops_daily'
    first=daily(folder/'data',folder/'outputs')
    second=daily(folder/'data',folder/'outputs')
    assert {k:v for k,v in first.items() if k!='elapsed_seconds'}=={k:v for k,v in second.items() if k!='elapsed_seconds'}
    book=load_workbook(folder/'outputs/运营日报.xlsx',data_only=True)
    assert book.sheetnames==['日报汇总','商品销量','库存预警','异常清单']
    overview=dict(book['日报汇总'].values)
    assert overview['销售额元']==1500
    assert overview['退款后销售额元']==1303
    assert book['异常清单'].max_row==13
    assert sum(row[2] for row in list(book['商品销量'].values)[1:])==10


def test_daily_missing_column_fails_and_marks_status(sample):
    folder=sample/'ops_daily'
    pd.DataFrame([dict(order_id='A')]).to_csv(folder/'data/orders_1.csv',index=False)
    with pytest.raises(ValueError,match='缺少字段'):
        daily(folder/'data',folder/'outputs')
    assert json.loads((folder/'outputs/metrics.json').read_text())['status']=='failed'


def test_duplicate_product_master_aborts(sample):
    folder=sample/'ops_daily'
    rows=read_table(folder/'data/products.xlsx',['sku'])
    write_workbook(folder/'data/products.xlsx',{'商品':(list(rows[0]),rows+[rows[0]])})
    with pytest.raises(ValueError,match='商品主数据'):
        daily(folder/'data',folder/'outputs')


@pytest.mark.parametrize('value',['NaN','Infinity','-1','1.001','','not-a-price'])
def test_bad_money_rejected(value):
    with pytest.raises(ValueError):amount(value)


@pytest.mark.parametrize('value',['1.5','-1','NaN',''])
def test_bad_inventory_rejected(value):
    with pytest.raises(ValueError):integer(value)


def test_catalog_validation():
    assert validate(dict(name='产品',sku='SKU',price='0.10',stock='0'))['price']==.1
    for row in [dict(name='',sku='A',price=1,stock=1),dict(name='a',sku='A',price=0,stock=1)]:
        with pytest.raises(ValueError):validate(row)


def test_formula_injection_is_text(tmp_path):
    write_workbook(tmp_path/'safe.xlsx',{'输入':(['name'],[dict(name='=HYPERLINK("bad")')])})
    cell=load_workbook(tmp_path/'safe.xlsx')['输入']['A2']
    assert cell.data_type=='s'
    assert cell.value.startswith("'=")


def test_real_browser_e2e_retry_timeout_and_idempotency(sample):
    folder=sample/'catalog_bot'
    db=folder/'state/test.sqlite'
    server=make_server(db,0,True)
    with sqlite3.connect(db) as connection:
        connection.execute('INSERT INTO products VALUES(?,?,?,?)',('PRE001','已有商品',1900,8))
    thread=threading.Thread(target=server.serve_forever,daemon=True)
    thread.start()
    url=f'http://127.0.0.1:{server.server_port}'
    try:
        result=robot(folder/'data/products.xlsx',folder/'outputs',url)
        assert result['counts']==dict(created=3,duplicate_input=1,invalid=3,failed=2,duplicate_existing=1)
        rows=json.loads((folder/'outputs/results.json').read_text(encoding='utf-8'))
        retry=next(r for r in rows if r['sku']=='RETRY-001')
        assert retry['attempts']==2 and retry['status']=='created'
        for sku in ['FAIL-001','SLOW-001']:
            row=next(r for r in rows if r['sku']==sku)
            assert row['attempts']==3
            assert (folder/'outputs'/row['screenshot']).stat().st_size>1000
        repeat=robot(folder/'data/products.xlsx',folder/'repeat',url)
        assert repeat['counts'].get('created',0)==0
        assert repeat['counts']['duplicate_existing']==4
        with sqlite3.connect(db) as connection:
            assert connection.execute('SELECT count(*) FROM products').fetchone()[0]==4
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
